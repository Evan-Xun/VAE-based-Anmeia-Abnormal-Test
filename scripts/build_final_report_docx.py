from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "report_assets"
OUTPUT = ROOT / "final_project_report_revised_detailed.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6" if edge in {"top", "left", "bottom", "right"} else "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "B8C2CC" if edge in {"top", "left", "bottom", "right"} else "D6DCE5")
        borders.append(el)
    tbl_pr.append(borders)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    title.font.size = Pt(28)
    title.font.bold = False
    title.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    for style_name, size, color in [
        ("Heading 1", 18, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 2", 15, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 3", 13, RGBColor(0x1F, 0x4D, 0x78)),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_paragraph(doc: Document, text: str, *, style: str | None = None, bold: bool = False, italic: bool = False,
                  align=None, color: RGBColor | None = None, size: Pt | None = None, space_after: int = 6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = size
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(f"• {text}")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(11)


def add_figure(doc: Document, image_name: str, caption: str, width: float = 6.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(ASSET_DIR / image_name), width=Inches(width))
    p.paragraph_format.space_after = Pt(4)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_table(doc: Document, rows: list[list[str]], column_widths: list[float]):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Inches(column_widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = value
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(10.5)
                    if i == 0:
                        run.bold = True
            if i == 0:
                set_cell_shading(cell, "EAF2FF")
    doc.add_paragraph()


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    configure_styles(doc)

    add_paragraph(doc, "基于 VAE 的 CBC 贫血异常检测项目报告", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        doc,
        "项目主题：VAE 基线实现、混合模型改进与端到端半监督方法评估",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        italic=True,
        color=RGBColor(0x66, 0x66, 0x66),
        size=Pt(14),
        space_after=2,
    )
    add_paragraph(
        doc,
        "文档说明：本报告依据项目实际实验流程与结果整理，面向课程 Final Report 提交。",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x66, 0x66, 0x66),
        size=Pt(12),
        space_after=14,
    )

    add_paragraph(doc, "摘要", style="Heading 1")
    add_paragraph(
        doc,
        "本项目围绕 CBC 贫血检测任务，首先实现了基于变分自编码器（VAE）的异常检测基线模型，并在结构化 v4 数据集上获得了较高的 AUROC，证明 VAE 基线在可分性较强的数据条件下具有可行性。随后，项目进一步构建了纯监督模型、松耦合混合模型以及端到端半监督 VAE，并在结构化数据、报告数据和融合数据三种条件下进行了系统比较。实验结果表明：VAE 基线对数据质量与分布一致性高度敏感；在 v4 数据集上，baseline VAE 可达到 AUROC 0.9223，而在 reports-only 数据集上下降到 0.5146；在融合数据集（latent_dim=4）上，hybrid + random_forest 获得当前最优 F1 0.8632。整体而言，本项目不仅验证了 VAE 基线的工程可行性，也揭示了不同数据条件下模型性能差异的来源，并为后续改进提供了明确方向。"
    )

    add_paragraph(doc, "1. 引言", style="Heading 1")
    add_paragraph(doc, "贫血是临床中常见的血液系统异常，早期识别对后续诊断和干预具有实际意义。全血细胞计数（CBC）是最常见、成本较低且获取方便的检查方式，因此基于 CBC 指标开展贫血检测具有明确应用价值。")
    add_paragraph(doc, "在传统监督学习设定下，若数据标签质量较高且特征完整，则分类模型通常可以取得较好结果。但在真实医疗数据场景中，常见问题包括标签粒度较粗、数据来源不一致、结构化特征有限，以及正常与异常之间边界不稳定。因此，本项目首先实现一个基于 VAE 的异常检测基线模型，希望通过学习正常样本分布识别贫血异常；随后进一步探索基于 VAE 的改进路线，包括松耦合混合模型与端到端半监督模型，并比较它们在不同数据条件下的表现。")
    add_paragraph(doc, "本项目主要回答以下问题：")
    add_bullet(doc, "单纯依赖 VAE 异常检测能否有效完成贫血检测？")
    add_bullet(doc, "将 VAE 衍生特征与监督分类器结合后，是否能稳定优于传统 VAE？")
    add_bullet(doc, "当数据集从结构化环境转向更复杂的报告数据时，模型性能会发生怎样的变化？")
    add_bullet(doc, "端到端半监督学习是否比“VAE 提特征 + 外部分类器”的松耦合方式更有效？")

    doc.add_page_break()
    add_paragraph(doc, "2. 方法与实验构建", style="Heading 1")
    add_paragraph(doc, "2.1 整体实验流程", style="Heading 2")
    add_paragraph(doc, "本项目的实验工作按照“数据准备 -> 基线实现 -> 改进模型构建 -> 多数据集评估 -> 结果分析”的顺序推进。首先完成数据清洗与 train / validation / test 划分；随后训练 baseline VAE，并在验证集上选择阈值；之后构建 supervised_only、vae_only、hybrid 与 end-to-end semi-supervised VAE；最后在测试集上进行比较。")
    add_figure(doc, "figure1_workflow.png", "图1  本项目整体实验流程图", 6.2)

    add_paragraph(doc, "2.2 数据集构建", style="Heading 2")
    add_paragraph(doc, "本项目在实验过程中使用了三种数据条件，它们分别代表了不同的数据复杂度。首先是结构化的 v4 数据集，其标签来自结构化诊断字段，特征较完整、来源相对统一，并且在健康与贫血之间具有较强可分性。其次是 reports-only 数据集，该数据集来自 CBC 文本报告解析后的结构化结果，仅保留 8 个核心指标：WBC、RBC、HGB、HCT、MCV、MCH、MCHC 和 PLT。最后，为检验模型在跨来源数据上的适应能力，项目构建了 8 指标融合数据集，将报告数据与 v4 中映射到相同 8 指标的样本合并，从而形成更复杂的混合型数据条件。")
    add_table(doc, [
        ["数据集", "来源", "特征数", "样本量", "特点"],
        ["v4", "结构化 CBC + Diagnosis", "14（主实验使用 8/14）", "清洗后约千级", "标签清晰、可分性强"],
        ["reports-only", "文本报告解析", "8", "962", "特征少、标签更粗、难度更高"],
        ["merged", "v4 + reports", "8", "约两千级", "跨来源分布更复杂"],
    ], [1.1, 1.7, 1.3, 1.0, 1.9])
    add_paragraph(doc, "表1  实验中使用的数据条件概览", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=RGBColor(0x66, 0x66, 0x66), size=Pt(10), space_after=10)
    add_paragraph(doc, "报告数据的清洗规则主要包括：剔除缺失值样本、剔除无法解析为数值的样本，以及剔除超出生理范围的样本。报告数据清洗结果为：original_rows = 1000，rows_after_cleaning = 962，dropped_rows = 38。所有实验统一采用分层切分：train / validation / test = 60 / 20 / 20，seed = 42。其中 baseline VAE 只使用训练集中的健康样本训练，阈值只在验证集选择，测试集只用于最终评估。")

    add_paragraph(doc, "2.3 模型设计", style="Heading 2")
    add_paragraph(doc, "本项目的方法部分按照“baseline VAE -> 松耦合 hybrid -> 端到端 semi-supervised VAE”的路线逐步展开。这样设计的原因是：首先需要验证基于正常样本分布建模的异常检测是否可行；在此基础上，再考察 VAE 学到的表示能否作为额外特征提升分类效果；最后进一步检验，把表示学习与分类学习放入同一优化框架后，是否能够减少目标错位并获得更强的任务相关表示。")
    add_paragraph(doc, "2.3.1 Baseline VAE 模型实现", style="Heading 3")
    add_paragraph(doc, "Baseline VAE 是一个面向表格型 CBC 指标的小型全连接网络。编码器采用 input_dim -> 16 -> 8 的两层 MLP，经 ReLU 激活后分别输出潜变量分布的均值 mu 和对数方差 log_var；解码器使用近似对称的 latent_dim -> 8 -> 16 -> input_dim 结构重构输入。之所以使用这样较浅的结构，是因为项目输入主要由少量数值型血常规指标组成，特征维度较低、语义明确，使用更深网络容易增加参数冗余与过拟合风险。默认情况下 latent_dim = 2，一方面有助于控制模型复杂度，另一方面也便于观察潜空间是否学到了紧凑的健康样本分布。")
    add_paragraph(doc, "在训练时，模型先由编码器得到 q(z|x) 的参数，再通过重参数化技巧完成采样：z = mu + sigma * epsilon，epsilon ~ N(0, I)。这样既保留了概率建模能力，又能进行稳定的梯度反传。损失函数写为 L = L_recon + beta * L_KL，其中 L_recon 使用均方误差衡量重构输入与原输入之间的差异，L_KL 用于约束潜变量分布接近标准正态。项目中 beta 默认设为 0.1，并采用 kl_warmup 机制，在前 50 个 epoch 内将 KL 权重从 0 线性增加到目标值。这一改动的动机，是避免训练初期 KL 项过强导致潜空间过早塌缩，让模型先建立稳定的正常样本重构能力，再逐步加入分布正则。")
    add_paragraph(doc, "在推理阶段，异常分数定义为 reconstruction error 与 latent KL contribution 之和，即 anomaly score = recon_error + kl_divergence。实现时直接使用 mu 进入解码器计算分数，而不是重复对 z 做随机采样，目的是减少采样噪声带来的预测波动，使阈值判别更稳定。baseline VAE 只在训练集中的健康样本上拟合，并且标准化器也仅由健康训练样本估计，从而尽量让模型专注于学习“正常 CBC 模式”。阈值选择放在验证集完成，测试集只用于最终评估，以避免信息泄漏。")
    add_paragraph(doc, "2.3.2 监督学习对照模型实现", style="Heading 3")
    add_paragraph(doc, "除了生成式路线之外，本项目还显式构建了监督学习对照模型，即 supervised_only。设置这一对照组的目的，是回答一个关键问题：在当前 CBC 数据条件下，是否必须引入 VAE 才能取得较好结果，还是仅依赖原始结构化指标就已经足够完成贫血分类。如果不加入这条监督学习基线，就很难判断后续 hybrid 或端到端模型的性能提升究竟来自方法设计，还是仅仅来自数据本身已经具有较强可分性。")
    add_paragraph(doc, "supervised_only 的输入直接使用标准化后的原始 CBC 指标，不包含任何 VAE 衍生特征。在实现上，项目统一复用了与 hybrid 相同的数据划分方式，即 train / validation / test = 60 / 20 / 20，并保持相同的随机种子与评价指标，从而保证不同模型路线之间具有可比性。与 baseline VAE 不同的是，监督学习模型在训练时直接使用训练集中的健康与贫血标签，因此它学习的是从原始特征到类别标签的判别边界，而不是正常样本分布本身。")
    add_paragraph(doc, "在分类器选择上，项目使用了两类具有代表性的监督模型：Logistic Regression 和 Random Forest。前者作为线性模型，适合衡量原始 CBC 指标是否已经可以被简单线性边界区分；后者可以进一步捕捉非线性关系与特征交互，用来检验原始指标中是否存在更复杂的判别结构。将这两个模型作为监督学习基线，有助于从线性可分性和非线性可分性两个角度理解数据难度。如果 supervised_only 已经表现很强，那么说明数据本身标签与特征关系较明确；如果它表现有限，则更能凸显引入 VAE 表示学习或生成式建模的必要性。")
    add_paragraph(doc, "此外，supervised_only 在本项目中不仅是一个简单的 baseline，更承担了解释实验结果的重要角色。例如，在结构化 v4 数据集上，监督学习模型接近饱和，说明该数据集本身更适合验证方法可行性；而在 reports-only 数据集上，监督模型虽优于 baseline VAE，但整体表现仍然有限，这进一步说明真实报告数据的任务难度更高。因此，监督学习对照组的加入，有助于让后续的 hybrid 与端到端结果具备更完整的参照系。")
    add_paragraph(doc, "2.3.3 松耦合混合模型实现", style="Heading 3")
    add_paragraph(doc, "为了改进 baseline VAE，本项目进一步构建松耦合 hybrid 模型。其核心想法并不是直接用 anomaly score 做最终分类，而是把 VAE 当作特征提取器。该部分系统比较三种特征使用方式：supervised_only（只使用原始 CBC 指标）、vae_only（只使用 VAE 衍生特征）和 hybrid（同时使用原始 CBC 指标与 VAE 衍生特征）。其中，VAE 衍生特征不仅包括 reconstruction error、KL divergence 和 anomaly score，还包括 latent mu、latent log_var，以及各个 CBC 指标的绝对重构误差与平方重构误差。这样设计的原因是：贫血相关信息可能分布在不同层次，单一 anomaly score 往往过于粗糙，而展开后的衍生特征能让监督分类器学习更细粒度的判别模式。")
    add_paragraph(doc, "在 hybrid 路线中，外部分类器统一使用 Logistic Regression 和 Random Forest。前者代表可解释性较强的线性基线，适合观察 VAE 特征是否带来稳定的线性增益；后者可以建模非线性关系和特征交互，用于检验 VAE 表示是否包含更复杂的信息补充。这里的 vae_only + classifier 主要作为消融实验，用来判断“仅依赖 VAE 表示”时是否已经具备足够判别力；若 hybrid 明显优于 supervised_only，则说明 VAE 特征确实为原始 CBC 指标提供了额外信息。")
    add_paragraph(doc, "2.3.4 端到端半监督 VAE 实现", style="Heading 3")
    add_paragraph(doc, "考虑到松耦合方法仍然存在目标错位问题，即 VAE 优化的是重构目标，而分类器优化的是判别目标，两者并非同时作用于编码器，本项目进一步实现了端到端半监督 VAE。该模型在 baseline VAE 的基础上增加一个分类头：输入 x 经 Encoder 得到 latent 表示后，同时进入 Decoder 与分类分支；分类头采用 latent_dim -> max(8, 2 * latent_dim) -> 1 的两层全连接结构，输出贫血类别 logits。总损失写为 L = L_recon + beta * L_KL + lambda * L_cls，其中 L_cls 使用二元交叉熵实现。分类分支直接使用 mu 而不是随机采样的 z 进行判别，是为了让监督信号更稳定地作用于潜空间中心，从而减少采样噪声对分类训练的干扰。")
    add_paragraph(doc, "这一改进的核心动机，是希望编码器学习到的潜在表示既能保留正常样本的结构信息，又能直接服务于贫血识别任务。与松耦合方法相比，端到端半监督结构让表示学习与分类学习在同一优化过程中共同约束编码器，理论上更有机会获得对任务更敏感的 latent representation。当然，这也意味着模型更依赖超参数平衡，例如 latent_dim、beta 与 classification_weight 的设置都会影响重构目标和分类目标之间的取舍，因此它是一个更具研究价值、但也更需要调参与验证的改进方向。")
    add_figure(doc, "figure2_model_routes.png", "图2  本项目比较的三条模型路线：Baseline、Hybrid 与 End-to-End", 6.2)

    add_paragraph(doc, "2.4 评价指标", style="Heading 2")
    add_paragraph(doc, "本项目主要使用 F1-score、AUROC、Precision、Recall、Specificity 与 AUPRC 评估模型表现。其中 F1-score 作为主指标，用于综合反映给定阈值下的分类效果；AUROC 用于评估模型整体排序能力；Recall 反映漏检情况；Specificity 反映误报情况。在医疗背景下，仅有较高的 AUROC 并不足以说明模型足够实用，还需要结合 Recall 与 Specificity 进行判断。")

    doc.add_page_break()
    add_paragraph(doc, "3. 实验结果", style="Heading 1")
    add_paragraph(doc, "3.1 Baseline VAE 在不同数据集上的表现", style="Heading 2")
    add_paragraph(doc, "baseline VAE 在三种数据条件下的表现差异非常明显。在结构化 v4 数据集上，baseline VAE 的 AUROC 达到 0.9223，F1 达到 0.9155；而在 reports-only 数据集上，AUROC 下降到 0.5146，几乎接近随机水平；在融合数据集（latent_dim=4）上，baseline 的 AUROC 回升到 0.6509，但仍明显弱于其他方法。")
    add_figure(doc, "figure3_baseline_cross_dataset.png", "图3  Baseline VAE 在不同数据条件下的 AUROC 与 F1 对比", 6.1)

    add_paragraph(doc, "3.2 v4 数据集结果", style="Heading 2")
    add_paragraph(doc, "v4 数据集的结果主要用于验证方法可行性。baseline VAE 已经表现良好，但纯监督与混合模型的结果进一步接近饱和。其中，supervised_only + random_forest 达到 AUROC 1.0000、F1 1.0000；hybrid + random_forest 达到 AUROC 0.9999、F1 0.9905；端到端半监督 VAE 也达到 AUROC 0.9868、F1 0.9712。由此可见，v4 数据集本身可分性过强，更适合说明模型能否工作，而不适合细分模型优劣。")
    add_table(doc, [
        ["模型", "AUROC", "F1", "Recall", "Specificity"],
        ["baseline VAE", "0.9223", "0.9155", "0.9874", "0.5970"],
        ["supervised_only + random_forest", "1.0000", "1.0000", "1.0000", "1.0000"],
        ["hybrid + random_forest", "0.9999", "0.9905", "0.9811", "1.0000"],
        ["end-to-end semi-supervised VAE", "0.9868", "0.9712", "0.9560", "0.9701"],
    ], [2.6, 1.0, 0.9, 1.1, 1.1])
    add_paragraph(doc, "表2  v4 数据集上的主要结果", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=RGBColor(0x66, 0x66, 0x66), size=Pt(10), space_after=10)
    add_figure(doc, "figure4_v4_comparison.png", "图4  v4 数据集上的模型分数对比", 6.1)

    add_paragraph(doc, "3.3 reports-only 数据集结果", style="Heading 2")
    add_paragraph(doc, "reports-only 数据集更接近真实困难场景。baseline VAE 在该数据集上的 AUROC 下降到 0.5146，说明其排序能力明显减弱；纯监督模型优于 baseline，其中 supervised_only + logistic_regression 获得该数据集主结果：AUROC 0.6202、F1 0.7013；hybrid 与 end-to-end 模型并未在这一数据集上体现稳定优势。")
    add_table(doc, [
        ["模型", "AUROC", "F1", "Recall", "Specificity"],
        ["baseline VAE", "0.5146", "0.6619", "0.9388", "0.0737"],
        ["supervised_only + logistic_regression", "0.6202", "0.7013", "0.8265", "0.4526"],
        ["hybrid + logistic_regression", "0.6332", "0.6862", "0.8367", "0.3789"],
        ["end-to-end semi-supervised VAE", "0.6339", "0.6507", "0.6939", "0.5474"],
    ], [2.6, 1.0, 0.9, 1.1, 1.1])
    add_paragraph(doc, "表3  reports-only 数据集上的主要结果", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=RGBColor(0x66, 0x66, 0x66), size=Pt(10), space_after=10)
    add_figure(doc, "figure5_report_comparison.png", "图5  reports-only 数据集上的模型分数对比", 6.1)
    add_paragraph(doc, "为了说明 baseline VAE 在报告数据上的训练与分数特性，图6给出了训练损失曲线，图7给出了测试集异常分数分布。可以看到，虽然模型训练过程是可收敛的，但健康样本与贫血样本的异常分数重叠较为严重，这也是 AUROC 明显下降的重要原因。")
    add_figure(doc, "figure7_training_loss.png", "图6  reports-only 数据集上 baseline VAE 的训练损失曲线", 6.0)
    add_figure(doc, "figure8_anomaly_distribution.png", "图7  reports-only 数据集上 baseline VAE 的异常分数分布", 6.0)

    add_paragraph(doc, "3.4 融合数据集结果（latent_dim = 4）", style="Heading 2")
    add_paragraph(doc, "在融合数据集上，baseline VAE 虽然较 reports-only 有一定回升，但仍明显弱于其他方法。纯监督、松耦合 hybrid 与端到端半监督三类方法都具有较高表现。其中，在 latent_dim = 4 的设定下，hybrid + random_forest 获得当前最佳结果：AUROC 0.8947、F1 0.8632；supervised_only + random_forest 的 AUROC 更高，为 0.9062，但 F1 略低于 hybrid；end-to-end semi-supervised VAE 获得 AUROC 0.8992、F1 0.8537，接近最优但未超越 hybrid + random_forest。")
    add_table(doc, [
        ["模型", "AUROC", "F1", "Recall", "Specificity"],
        ["baseline VAE", "0.6509", "0.7500", "0.9222", "0.1481"],
        ["supervised_only + random_forest", "0.9062", "0.8552", "0.9650", "0.5370"],
        ["hybrid + random_forest", "0.8947", "0.8632", "0.9572", "0.5864"],
        ["end-to-end semi-supervised VAE", "0.8992", "0.8537", "0.9650", "0.5309"],
    ], [2.6, 1.0, 0.9, 1.1, 1.1])
    add_paragraph(doc, "表4  融合数据集（latent_dim = 4）上的主要结果", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=RGBColor(0x66, 0x66, 0x66), size=Pt(10), space_after=10)
    add_figure(doc, "figure6_merged_comparison.png", "图8  融合数据集（latent_dim = 4）上的模型分数对比", 6.1)

    doc.add_page_break()
    add_paragraph(doc, "4. Conclusion", style="Heading 1")
    add_paragraph(doc, "4.1 Results Analysis", style="Heading 2")
    add_paragraph(doc, "本项目从实现 VAE 基线出发，逐步扩展到纯监督、松耦合混合以及端到端半监督方法，并在不同 CBC 数据条件下进行了系统比较。综合结果可以得出以下结论：")
    add_bullet(doc, "VAE 基线在结构化、可分性较强的 v4 数据集上表现良好，AUROC 达到 0.9223，说明其作为贫血异常检测基线是可行的。")
    add_bullet(doc, "v4 数据集本身可分性过强，不足以充分区分改进方法优劣。")
    add_bullet(doc, "reports-only 数据集更接近真实困难场景，baseline VAE 在该数据上下降明显；正式结果中，supervised_only + logistic_regression 与 hybrid + logistic_regression 的 AUROC 分别为 0.6202 和 0.6332。")
    add_bullet(doc, "hybrid 模型具有一定价值，但优势不稳定；在融合数据集上表现最好，在报告数据上则未稳定优于纯监督模型。")
    add_bullet(doc, "端到端半监督 VAE 是合理的研究方向，但当前实现仍需继续调参与结构优化。")
    add_paragraph(doc, "从研究角度看，本项目最重要的发现并不是某个改进模型必然更强，而是模型性能高度依赖于数据条件。结构化程度高、标签纯度高、特征完整的数据集，更容易让 VAE 和监督模型都取得很高结果；而在更真实、更复杂的数据条件下，模型优劣会发生明显变化。")
    add_paragraph(doc, "本项目的意义在于：完成了 VAE 基线模型的工程实现与系统评估；比较了纯 VAE、纯监督、混合和端到端半监督四类方法；揭示了不同 CBC 数据集之间任务难度的显著差异；并明确了当前方法的适用边界，为后续改进提供了方向。")

    add_paragraph(doc, "4.2 Future Improvements", style="Heading 2")
    add_paragraph(doc, "后续工作可从以下方向推进：一是增加更多 CBC 特征，如 RDW、MPV、分类百分比和绝对计数；二是进一步改善报告数据标签质量；三是开展多 seed 或交叉验证实验，报告均值与标准差；四是继续优化端到端半监督结构中的 latent_dim 与分类损失权重；五是尝试 Deep SVDD-VAE、条件 VAE 或其他更适合异常检测的生成式方法。")

    add_paragraph(doc, "4.3 Future Study Topics", style="Heading 2")
    add_bullet(doc, "医疗表格数据中的半监督异常检测。")
    add_bullet(doc, "跨数据源 CBC 检测模型的鲁棒性研究。")
    add_bullet(doc, "噪声标签条件下的贫血检测。")
    add_bullet(doc, "结构化 CBC 数据与文本报告联合建模。")
    add_bullet(doc, "面向临床使用的低误报、高召回贫血筛查模型设计。")

    add_paragraph(doc, "5. Reference", style="Heading 1")
    refs = [
        "D. P. Kingma and M. Welling, Auto-Encoding Variational Bayes, ICLR, 2014.",
        "J. An and S. Cho, Variational Autoencoder based Anomaly Detection using Reconstruction Probability, Special Lecture on IE, 2015.",
        "L. Ruff, R. A. Vandermeulen, N. Görnitz, et al., Deep One-Class Classification, ICML, 2018.",
        "F. Pedregosa, G. Varoquaux, A. Gramfort, et al., Scikit-learn: Machine Learning in Python, JMLR, 2011.",
        "A. Paszke, S. Gross, F. Massa, et al., PyTorch: An Imperative Style, High-Performance Deep Learning Library, NeurIPS, 2019.",
    ]
    for idx, ref in enumerate(refs, start=1):
        add_paragraph(doc, f"[{idx}] {ref}", space_after=2)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
